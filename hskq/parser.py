"""Đọc bảng từ vựng HSK trong file .docx -> danh sách thẻ (Card)."""
from __future__ import annotations

import re
from dataclasses import dataclass, asdict, field
from pathlib import Path

import docx

# Nhận diện cột theo tiêu đề (không phụ thuộc thứ tự cột)
COLS = {
    "hanzi":   ["chữ hán", "汉字", "hán tự", "từ vựng", "生词"],
    "pinyin":  ["pinyin", "phiên âm", "拼音"],
    "meaning": ["giải nghĩa", "nghĩa", "từ loại"],
    "example": ["ví dụ", "例句"],
    "related": ["liên quan", "đồng nghĩa", "bộ thủ"],
}


@dataclass
class Card:
    index: int
    hanzi: str
    pinyin: str = ""
    meaning: str = ""
    example: str = ""
    related: str = ""
    term: str = ""          # mặt trước
    definition: str = ""    # mặt sau
    keyword: str = ""       # từ khóa tìm hình
    image: str = ""         # đường dẫn hình (tương đối với thư mục bài)

    def to_dict(self):
        return asdict(self)


@dataclass
class Lesson:
    source: str
    number: str = ""
    title: str = ""
    cards: list[Card] = field(default_factory=list)

    @property
    def slug(self) -> str:
        return f"Bai_{self.number}" if self.number else Path(self.source).stem


def _clean(s: str) -> str:
    s = s.replace("\t", " ").replace(" ", " ").replace("➔", "→")
    lines = [re.sub(r"\s+", " ", l).strip() for l in s.splitlines()]
    return "\n".join(l for l in lines if l)


def _map_columns(header: list[str]):
    h = [c.lower() for c in header]
    found: dict[str, int] = {}
    for key, kws in COLS.items():
        for i, col in enumerate(h):
            if i not in found.values() and any(k in col for k in kws):
                found[key] = i
                break
    if "hanzi" not in found or "meaning" not in found:
        return None
    return found


def _cell_text(cell) -> str:
    return _clean("\n".join(p.text for p in cell.paragraphs))


def guess_keyword(meaning: str) -> str:
    """Lấy nghĩa tiếng Việt đầu tiên làm từ khóa tìm hình: '→ Bệnh nhân, người bệnh' -> 'Bệnh nhân'."""
    for line in meaning.splitlines():
        if "→" in line:
            line = line.split("→", 1)[1]
        elif line.startswith("("):
            continue
        line = re.sub(r"\(.*?\)", "", line)
        kw = re.split(r"[,;/]", line)[0].strip(" .")
        if kw:
            return kw
    return ""


def _lesson_meta(d) -> tuple[str, str]:
    text = "\n".join(p.text for p in d.paragraphs[:15])
    m = re.search(r"Bài\s*(\d+)\s*\(([^)]+)\)", text, re.I)
    if m:
        return m.group(1), m.group(2).strip()
    m = re.search(r"BÀI\s*(\d+)\s*[:：]\s*(.+)", text, re.I)
    if m:
        return m.group(1), m.group(2).strip()
    return "", ""


def parse_docx(path: str | Path, include_related: bool = True) -> Lesson:
    path = Path(path)
    d = docx.Document(str(path))
    num, title = _lesson_meta(d)
    if not num:
        m = re.search(r"(\d+)", path.stem)
        num = m.group(1) if m else ""
    lesson = Lesson(source=str(path), number=num, title=title)

    for t in d.tables:
        rows = [[_cell_text(c) for c in r.cells] for r in t.rows]
        if not rows:
            continue
        cols = _map_columns(rows[0])
        if not cols:
            continue
        for r in rows[1:]:
            def get(k):
                return r[cols[k]] if k in cols and cols[k] < len(r) else ""
            hanzi = get("hanzi")
            if not hanzi or _map_columns(r):   # dòng trống / tiêu đề lặp lại
                continue
            c = Card(index=len(lesson.cards) + 1, hanzi=hanzi, pinyin=get("pinyin"),
                     meaning=get("meaning"), example=get("example"), related=get("related"))
            c.term = c.hanzi                                  # mặt trước: chỉ chữ Hán
            parts = [c.pinyin, c.meaning]                     # mặt sau: pinyin trước
            if c.example:
                parts.append("Ví dụ: " + c.example)
            if include_related and c.related:
                parts.append(c.related)
            c.definition = "\n".join(p for p in parts if p)
            c.keyword = guess_keyword(c.meaning)
            lesson.cards.append(c)
    return lesson
